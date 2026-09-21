#!/usr/bin/env python3
"""Add ``rna_type`` + ``protein_type`` columns to the test-results .xlsx.

Merges the two classification files keyed by ``sample_id``:

* ``classification_summary.csv`` — the Rfam (cmsearch) / Pfam (hmmscan)
  result. **Takes priority.**
* ``pdb_entity_descriptions.csv`` — the RCSB ``_entity.pdbx_description``
  fallback (with the keyword ``rna_type_guess``). Used only when the
  Rfam/Pfam side is ``unclassified`` / empty.

Per sample::

    rna_type     = rna_rfam_family   if classified else rna_type_guess
    protein_type = protein_pfam_domain if classified else protein_description

The two columns are appended to the existing sheet (existing columns and
their values are untouched). Re-running is idempotent: if the columns
already exist they are overwritten in place rather than duplicated.

Usage
-----
::

    python scripts/riboseer/merge_classification_into_results.py \\
        --results data/batch_test_v7/test_results.xlsx \\
        --rfam-csv data/batch_test_v7/classification/classification_summary.csv \\
        --pdb-entity-csv data/batch_test_v7/classification/pdb_entity_descriptions.csv
        # writes back in place; pass --output to write a copy instead.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path
from typing import Optional

RNA_TYPE_COL = "rna_type"
PROTEIN_TYPE_COL = "protein_type"
UNCLASSIFIED = "unclassified"
UNKNOWN = "unknown"


def _is_classified(v: Optional[str]) -> bool:
    return bool(v) and v.strip().lower() not in (UNCLASSIFIED, "", "none", "na")


def load_csv(path: Path) -> dict[str, dict]:
    """``{sample_id: row_dict}`` from a CSV keyed on ``sample_id``."""
    out: dict[str, dict] = {}
    if not path.is_file():
        return out
    with path.open(encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            sid = (row.get("sample_id") or "").strip()
            if sid:
                out[sid] = row
    return out


def merge_types(sid: str, rfam: dict[str, dict], pdb: dict[str, dict]
                ) -> tuple[str, str]:
    """``(rna_type, protein_type)`` for one sample — Rfam/Pfam first, then
    the PDB-entity fallback, then ``unknown``."""
    rf = rfam.get(sid) or {}
    pe = pdb.get(sid) or {}

    rfam_fam = rf.get("rna_rfam_family")
    if _is_classified(rfam_fam):
        rna_type = rfam_fam.strip()
    else:
        rna_type = (pe.get("rna_type_guess") or "").strip() or UNKNOWN

    pfam_dom = rf.get("protein_pfam_domain")
    if _is_classified(pfam_dom):
        protein_type = pfam_dom.strip()
    else:
        protein_type = (pe.get("protein_description") or "").strip() or UNKNOWN

    return rna_type, protein_type


def patch_xlsx(results: Path, rfam: dict[str, dict], pdb: dict[str, dict],
               out: Path) -> tuple[int, int, int]:
    """Append / refresh the two type columns. Returns
    ``(n_rows, n_rna_from_fallback, n_protein_from_fallback)``."""
    from openpyxl import load_workbook

    wb = load_workbook(str(results))
    ws = wb["test_results"] if "test_results" in wb.sheetnames else wb.active

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if "sample_id" not in header:
        raise ValueError("results sheet has no 'sample_id' column")
    sid_col = header.index("sample_id") + 1               # 1-based

    # Reuse existing type columns (idempotent) or append new ones.
    def col_for(name: str) -> int:
        if name in header:
            return header.index(name) + 1
        ws.cell(row=1, column=ws.max_column + 1, value=name)
        header.append(name)
        return ws.max_column

    rna_col = col_for(RNA_TYPE_COL)
    prot_col = col_for(PROTEIN_TYPE_COL)

    n_rows = rna_fb = prot_fb = 0
    for r in range(2, ws.max_row + 1):
        sid = ws.cell(row=r, column=sid_col).value
        if sid is None:
            continue
        sid = str(sid).strip()
        rna_type, protein_type = merge_types(sid, rfam, pdb)
        # Track fallback usage for the report.
        rf = rfam.get(sid) or {}
        if not _is_classified(rf.get("rna_rfam_family")):
            rna_fb += 1
        if not _is_classified(rf.get("protein_pfam_domain")):
            prot_fb += 1
        ws.cell(row=r, column=rna_col, value=rna_type)
        ws.cell(row=r, column=prot_col, value=protein_type)
        n_rows += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return n_rows, rna_fb, prot_fb


def main(argv: Optional[list[str]] = None) -> int:
    p = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--results", type=Path, required=True,
                   help="test_results.xlsx to enrich.")
    p.add_argument("--rfam-csv", type=Path, required=True,
                   help="classification_summary.csv (cmsearch/hmmscan).")
    p.add_argument("--pdb-entity-csv", type=Path, required=True,
                   help="pdb_entity_descriptions.csv (RCSB fallback).")
    p.add_argument("--output", type=Path, default=None,
                   help="output path; defaults to --results (in place).")
    args = p.parse_args(argv)

    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    if not args.results.is_file():
        print(f"ERROR: results not found: {args.results}", file=sys.stderr)
        return 1
    rfam = load_csv(args.rfam_csv)
    pdb = load_csv(args.pdb_entity_csv)
    if not rfam:
        print(f"WARNING: no rows in {args.rfam_csv}", file=sys.stderr)
    if not pdb:
        print(f"WARNING: no rows in {args.pdb_entity_csv}", file=sys.stderr)
    print(f"loaded rfam/pfam: {len(rfam)}  pdb-entity: {len(pdb)}")

    out = args.output or args.results
    try:
        n, rna_fb, prot_fb = patch_xlsx(args.results, rfam, pdb, out)
    except ImportError:
        print("ERROR: openpyxl required (pip install openpyxl)",
              file=sys.stderr)
        return 1
    print(f"wrote {out}")
    print(f"  {n} rows enriched with '{RNA_TYPE_COL}' + '{PROTEIN_TYPE_COL}'")
    print(f"  RNA: {n - rna_fb} from Rfam, {rna_fb} from PDB-entity fallback")
    print(f"  protein: {n - prot_fb} from Pfam, {prot_fb} from PDB-entity "
          f"fallback")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
