"""Mock tests for scripts/riboseer/export_test_results_excel.py.

Row construction (meta + per-method Pearson + RiboSeer extras) runs
locally without sklearn/openpyxl; the actual .xlsx round-trip is guarded
by skipUnless(openpyxl).
"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(REPO_ROOT / "src"), str(REPO_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from scripts.riboseer import export_test_results_excel as ex  # noqa: E402
from step5_fusion.prediction_io import load_predictions_dir  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False


def _pred(tool_id, per_res, success=True):
    return {
        "tool_id": tool_id, "success": success,
        "binding_protein_residues": [],
        "per_residue_pae_score": {str(r): v for r, v in per_res.items()},
    }


def _setup(td):
    root = Path(td)
    samples = root / "data" / "samples"
    step4 = root / "step4"
    preds = root / "preds"
    for d in (samples, step4, preds):
        d.mkdir(parents=True, exist_ok=True)
    sid = "4n0t_A_B"
    (samples / f"{sid}.json").write_text(json.dumps({
        "sample_id": sid, "source_pdb": "4n0t",
        "protein": {"chain_id": "A", "length": 6, "sequence": "MAKLPQ",
                    "resolved_residues": [1, 2, 3, 4, 5, 6]},
        "rna": {"chain_id": "B", "length": 8, "sequence": "ACGUACGU"},
        "interaction": {"binding_protein_residues": [1, 2, 3]},
    }), encoding="utf-8")
    (step4 / f"{sid}.jsonl").write_text(json.dumps({
        "sample_id": sid, "predictions": [
            _pred("boltz2", {1: 0.8, 2: 0.6, 3: 0.5, 4: 0.55, 5: 0.2, 6: 0.1}),
            _pred("equipnas", {1: 0.7, 2: 0.55, 3: 0.5,
                               4: 0.45, 5: 0.3, 6: 0.1}),
        ]}), encoding="utf-8")
    (preds / f"{sid}.json").write_text(json.dumps({
        "residue_ids": [1, 2, 3, 4, 5, 6],
        "probabilities": [0.97, 0.95, 0.93, 0.05, 0.04, 0.02],
    }), encoding="utf-8")
    return root, samples.parent, step4, preds, sid


class TestColumns(unittest.TestCase):
    def test_column_count_and_order(self):
        # 9 meta + 1 RiboSeer Pearson + 15 baselines + 6 extras = 31
        self.assertEqual(len(ex.COLUMNS), 9 + 1 + 15 + 6)
        self.assertEqual(ex.COLUMNS[0], "sample_id")
        self.assertEqual(ex.COLUMNS[9], "RiboSeer_PearsonR")
        self.assertIn("Boltz2_PearsonR", ex.COLUMNS)
        self.assertIn("BindUP_PearsonR", ex.COLUMNS)
        self.assertIn("RiboSeer_MCC", ex.COLUMNS)
        # baselines only get Pearson, no Spearman/MCC etc.
        self.assertNotIn("Boltz2_MCC", ex.COLUMNS)


class TestBuildRow(unittest.TestCase):
    def test_build_rows_no_sklearn(self):
        with tempfile.TemporaryDirectory() as td:
            root, data_dir, step4, preds, sid = _setup(td)
            fs = load_predictions_dir(preds)
            rows, skipped = ex.build_rows(
                [sid], data_dir=data_dir, step4_dir=step4,
                fs_preds=fs, sklearn_ok=False)
            self.assertEqual(skipped, 0)
            self.assertEqual(len(rows), 1)
            r = rows[0]
            # meta
            self.assertEqual(r["sample_id"], sid)
            self.assertEqual(r["pdb_code"], "4n0t")
            self.assertEqual(r["protein_chain"], "A")
            self.assertEqual(r["rna_chain"], "B")
            self.assertEqual(r["protein_length"], 6)
            self.assertEqual(r["rna_length"], 8)
            self.assertEqual(r["protein_sequence"], "MAKLPQ")
            self.assertEqual(r["rna_sequence"], "ACGUACGU")
            self.assertEqual(r["n_binding_residues"], 3)
            # Pearson present for the two tools, NaN for the rest
            self.assertNotEqual(r["Boltz2_PearsonR"], ex.NAN)
            self.assertNotEqual(r["EquiPNAS_PearsonR"], ex.NAN)
            self.assertEqual(r["HDOCK_PearsonR"], ex.NAN)
            self.assertEqual(r["NucleicNet_PearsonR"], ex.NAN)
            self.assertNotEqual(r["RiboSeer_PearsonR"], ex.NAN)
            # Spearman/R2 always available (no sklearn needed)
            self.assertNotEqual(r["RiboSeer_SpearmanR"], ex.NAN)
            self.assertNotEqual(r["RiboSeer_R2"], ex.NAN)
            # sklearn-only metrics -> NaN when sklearn_ok=False
            self.assertEqual(r["RiboSeer_AUROC"], ex.NAN)
            self.assertEqual(r["RiboSeer_MCC"], ex.NAN)
            # every declared column is present
            for c in ex.COLUMNS:
                self.assertIn(c, r)

    def test_rows_sorted_by_sample_id(self):
        with tempfile.TemporaryDirectory() as td:
            root, data_dir, step4, preds, sid = _setup(td)
            # add a second sample with an earlier id
            (Path(data_dir) / "samples" / "1aaa_A_B.json").write_text(
                json.dumps({
                    "sample_id": "1aaa_A_B", "source_pdb": "1aaa",
                    "protein": {"chain_id": "A", "length": 4,
                                "sequence": "MAKL",
                                "resolved_residues": [1, 2, 3, 4]},
                    "rna": {"chain_id": "B", "length": 4, "sequence": "ACGU"},
                    "interaction": {"binding_protein_residues": [1, 2]},
                }), encoding="utf-8")
            (step4 / "1aaa_A_B.jsonl").write_text(json.dumps({
                "predictions": [
                    _pred("boltz2", {1: 0.9, 2: 0.6, 3: 0.4, 4: 0.5})]}),
                encoding="utf-8")
            (preds / "1aaa_A_B.json").write_text(json.dumps({
                "residue_ids": [1, 2, 3, 4],
                "probabilities": [0.9, 0.8, 0.1, 0.05]}), encoding="utf-8")
            fs = load_predictions_dir(preds)
            rows, _ = ex.build_rows(
                [sid, "1aaa_A_B"], data_dir=data_dir, step4_dir=step4,
                fs_preds=fs, sklearn_ok=False)
            self.assertEqual([r["sample_id"] for r in rows],
                             ["1aaa_A_B", "4n0t_A_B"])


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed (server-only)")
class TestWriteXlsx(unittest.TestCase):
    def test_roundtrip(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as td:
            root, data_dir, step4, preds, sid = _setup(td)
            fs = load_predictions_dir(preds)
            rows, _ = ex.build_rows(
                [sid], data_dir=data_dir, step4_dir=step4,
                fs_preds=fs, sklearn_ok=False)
            out = root / "out.xlsx"
            ex.write_xlsx(rows, out)
            self.assertTrue(out.is_file())
            wb = openpyxl.load_workbook(out)
            ws = wb.active
            header = [c.value for c in ws[1]]
            self.assertEqual(header, ex.COLUMNS)
            self.assertEqual(ws.freeze_panes, "A2")
            # one data row
            self.assertEqual(ws.max_row, 2)
            data = {h: ws.cell(row=2, column=i + 1).value
                    for i, h in enumerate(header)}
            self.assertEqual(data["sample_id"], sid)
            self.assertEqual(data["protein_sequence"], "MAKLPQ")
            self.assertEqual(data["HDOCK_PearsonR"], "NaN")


if __name__ == "__main__":
    unittest.main()
