"""Tests for scripts/riboseer/merge_classification_into_results.py.

Covers the Rfam/Pfam-priority + PDB-entity-fallback merge rule, appending
the two columns to a real .xlsx, and idempotent re-runs.
"""
import sys
import tempfile
import unittest
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
for _p in (str(REPO_ROOT / "src"), str(REPO_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from scripts.riboseer import merge_classification_into_results as mc  # noqa: E402


def _write_csv(path, header, rows):
    import csv
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def _make_xlsx(path, header, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "test_results"
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(str(path))


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    ws = wb["test_results"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(header, r)))
    return header, rows


class TestMergeRule(unittest.TestCase):
    def test_rfam_priority(self):
        rfam = {"s1": {"sample_id": "s1", "rna_rfam_family": "tmRNA",
                       "protein_pfam_domain": "SmpB"}}
        pdb = {"s1": {"sample_id": "s1", "rna_type_guess": "tRNA/tmRNA",
                      "protein_description": "SsrA-binding protein"}}
        self.assertEqual(mc.merge_types("s1", rfam, pdb), ("tmRNA", "SmpB"))

    def test_fallback_when_unclassified(self):
        rfam = {"s2": {"sample_id": "s2", "rna_rfam_family": "unclassified",
                       "protein_pfam_domain": "Ribosomal_L1"}}
        pdb = {"s2": {"sample_id": "s2",
                      "rna_type_guess": "synthetic/generic-oligo",
                      "protein_description": "50S ribosomal protein L1"}}
        # RNA falls back (rfam unclassified); protein keeps Pfam.
        self.assertEqual(mc.merge_types("s2", rfam, pdb),
                         ("synthetic/generic-oligo", "Ribosomal_L1"))

    def test_protein_fallback(self):
        rfam = {"s3": {"sample_id": "s3", "rna_rfam_family": "tRNA",
                       "protein_pfam_domain": ""}}
        pdb = {"s3": {"sample_id": "s3", "rna_type_guess": "x",
                      "protein_description": "Cas9"}}
        self.assertEqual(mc.merge_types("s3", rfam, pdb), ("tRNA", "Cas9"))

    def test_unknown_when_both_missing(self):
        self.assertEqual(mc.merge_types("zz", {}, {}), ("unknown", "unknown"))


class TestPatchXlsx(unittest.TestCase):
    def _setup(self, td):
        xlsx = td / "results.xlsx"
        _make_xlsx(xlsx, ["sample_id", "RiboSeer_PearsonR"],
                   [["s1", 0.6], ["s2", 0.5]])
        rfam_csv = td / "rfam.csv"
        _write_csv(rfam_csv,
                   ["sample_id", "rna_rfam_family", "protein_pfam_domain"],
                   [["s1", "tmRNA", "SmpB"],
                    ["s2", "unclassified", "Ribosomal_L1"]])
        pdb_csv = td / "pdb.csv"
        _write_csv(pdb_csv,
                   ["sample_id", "rna_type_guess", "protein_description"],
                   [["s1", "tRNA/tmRNA", "SsrA"],
                    ["s2", "synthetic/generic-oligo", "L1"]])
        return xlsx, rfam_csv, pdb_csv

    def test_appends_two_columns(self):
        with tempfile.TemporaryDirectory() as t:
            td = Path(t)
            xlsx, rfam_csv, pdb_csv = self._setup(td)
            n, rna_fb, prot_fb = mc.patch_xlsx(
                xlsx, mc.load_csv(rfam_csv), mc.load_csv(pdb_csv), xlsx)
            self.assertEqual(n, 2)
            self.assertEqual(rna_fb, 1)      # s2 RNA from fallback
            self.assertEqual(prot_fb, 0)
            header, rows = _read_xlsx(xlsx)
            self.assertIn("rna_type", header)
            self.assertIn("protein_type", header)
            # existing column preserved.
            self.assertIn("RiboSeer_PearsonR", header)
            by = {r["sample_id"]: r for r in rows}
            self.assertEqual(by["s1"]["rna_type"], "tmRNA")
            self.assertEqual(by["s2"]["rna_type"], "synthetic/generic-oligo")
            self.assertEqual(by["s2"]["protein_type"], "Ribosomal_L1")

    def test_idempotent_rerun(self):
        with tempfile.TemporaryDirectory() as t:
            td = Path(t)
            xlsx, rfam_csv, pdb_csv = self._setup(td)
            args = (mc.load_csv(rfam_csv), mc.load_csv(pdb_csv))
            mc.patch_xlsx(xlsx, *args, xlsx)
            mc.patch_xlsx(xlsx, *args, xlsx)         # run twice
            header, _ = _read_xlsx(xlsx)
            # columns appear exactly once.
            self.assertEqual(header.count("rna_type"), 1)
            self.assertEqual(header.count("protein_type"), 1)


if __name__ == "__main__":
    unittest.main()
